import hashlib
import io
import json
import sqlite3
import zipfile
from datetime import datetime, timezone
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook = None
    load_workbook = None

SHEETS = {
    "cashier": """
        SELECT transaction_id,created_at,device_id,member_id,credential_id,account_id,
               'DEPOSIT' AS transaction_type,amount,balance_before,balance_after,status
        FROM cashier_deposits
        ORDER BY created_at
    """,
    "vending": """
        SELECT transaction_id,created_at,machine_id,device_id,credential_id,account_id,
               'PURCHASE' AS transaction_type,product_id,amount,payment_status,
               dispense_status,status
        FROM vending_transactions v
        LEFT JOIN vending_machines m ON m.machine_id=v.machine_id
        ORDER BY created_at
    """,
    "payments": """
        SELECT transaction_id,created_at,transaction_source,device_id,credential_id,
               account_id,method,provider,original_amount,discount_amount,final_amount,
               amount,status,idempotency_key
        FROM payment_transactions
        ORDER BY created_at
    """,
    "ledger": """
        SELECT id,created_at,account_id,direction,amount,reference
        FROM financial_ledger
        ORDER BY created_at
    """,
    "audit": """
        SELECT id,created_at,event_type,entity_type,entity_id,detail
        FROM audit_events
        ORDER BY created_at
    """,
}

def _rows(db, sql):
    with db.connect() as c:
        cur=c.execute(sql)
        cols=[d[0] for d in cur.description]
        return cols,[tuple(r) for r in cur.fetchall()]

def export_excel(db, target, start=None, end=None):
    if Workbook is None:
        raise RuntimeError("openpyxl_required")
    target=Path(target)
    target.parent.mkdir(parents=True,exist_ok=True)
    wb=Workbook()
    wb.remove(wb.active)
    for name,sql in SHEETS.items():
        if start and "WHERE" not in sql.upper():
            sql=sql.replace("ORDER BY","WHERE created_at>=? AND created_at<? ORDER BY")
            args=(start,end)
        else:
            args=()
        with db.connect() as c:
            cur=c.execute(sql,args)
            headers=[d[0] for d in cur.description]
            rows=cur.fetchall()
        ws=wb.create_sheet(name[:31])
        ws.append(headers)
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes="A2"
        ws.auto_filter.ref=ws.dimensions
    wb.save(target)
    return target

def create_native_backup(db, target, include_excel=True):
    target=Path(target)
    target.parent.mkdir(parents=True,exist_ok=True)
    with db.connect() as c:
        raw=c.execute("SELECT name,sql FROM sqlite_master WHERE type='table' ORDER BY name").fetchall()
        dump="\n".join(c.iterdump()).encode()
    manifest={
        "format":"nvm-native-backup-v1",
        "created_at":datetime.now(timezone.utc).isoformat(),
        "retention_months":12,
        "tables":[r[0] for r in raw],
    }
    with zipfile.ZipFile(target,"w",zipfile.ZIP_DEFLATED) as z:
        z.writestr("database.sql",dump)
        z.writestr("manifest.json",json.dumps(manifest,indent=2))
        if include_excel:
            if Workbook is None:
                raise RuntimeError("openpyxl_required")
            buf=target.with_suffix(".xlsx")
            export_excel(db,buf)
            z.write(buf,"excel/nvm-export.xlsx")
            buf.unlink(missing_ok=True)
        digest=hashlib.sha256(dump).hexdigest()
        z.writestr("checksum.sha256",digest+"  database.sql\n")
    return target

def restore_native_backup(db, archive):
    archive=Path(archive)
    with zipfile.ZipFile(archive) as z:
        manifest=json.loads(z.read("manifest.json"))
        dump=z.read("database.sql")
        expected=z.read("checksum.sha256").decode().split()[0]
        if hashlib.sha256(dump).hexdigest()!=expected:
            raise ValueError("backup_checksum_mismatch")
        sql=dump.decode()
    # Validate the SQL dump in an isolated temporary SQLite database first.
    # Only replace the live database after the dump has been proven loadable.
    tmp = archive.with_suffix(".restore.db")
    tmp.unlink(missing_ok=True)
    probe = None
    try:
        probe = sqlite3.connect(tmp)
        probe.execute("PRAGMA foreign_keys=ON")
        probe.executescript(sql)
        probe.commit()
        probe.close()
        probe = None

        live = Path(db.path)
        live.parent.mkdir(parents=True, exist_ok=True)
        tmp.replace(live)
    except Exception:
        if probe is not None:
            probe.close()
        tmp.unlink(missing_ok=True)
        raise
    return manifest

def import_excel(db, source):
    if load_workbook is None:
        raise RuntimeError("openpyxl_required")
    wb=load_workbook(source,data_only=True,read_only=True)
    imported={}
    with db.connect() as c:
        if "cashier" in wb.sheetnames:
            ws=wb["cashier"]
            headers=[x.value for x in next(ws.iter_rows())]
            idx={h:i for i,h in enumerate(headers)}
            for row in ws.iter_rows(min_row=2,values_only=True):
                tid=row[idx["transaction_id"]]
                if not tid: continue
                c.execute("""INSERT OR IGNORE INTO cashier_deposits
                    (transaction_id,device_id,credential_id,member_id,account_id,amount,status,
                     balance_before,balance_after,created_at,completed_at)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (tid,row[idx["device_id"]],row[idx["credential_id"]],row[idx["member_id"]],
                     row[idx["account_id"]],int(row[idx["amount"]]),row[idx["status"]],
                     row[idx["balance_before"]],row[idx["balance_after"]],row[idx["created_at"]],row[idx["created_at"]]))
            imported["cashier"]=c.total_changes
        c.commit()
    return imported


def export_rows_excel(rows, target):
    if Workbook is None:
        raise RuntimeError("openpyxl_required")
    target=Path(target)
    target.parent.mkdir(parents=True,exist_ok=True)
    wb=Workbook()
    ws=wb.active
    ws.title="audit"
    if rows:
        headers=list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
    else:
        ws.append(["transaction_id","created_at","source","status","amount"])
    ws.freeze_panes="A2"
    ws.auto_filter.ref=ws.dimensions
    wb.save(target)
    return target
